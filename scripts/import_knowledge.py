#!/usr/bin/env python3
import argparse, gzip, hashlib, json, re, shutil, subprocess, zipfile
from pathlib import Path

TEXT_EXT = {".txt",".md",".csv",".tsv",".json",".jsonl",".xml",".html",".htm",".ini",".cfg",".conf",".yaml",".yml",".log",".sql",".rtf",".inf"}
IMAGE_EXT = {".jpg",".jpeg",".png",".gif",".bmp",".tif",".tiff",".webp"}
BINARY_TEXT_EXT = {".cxt",".cst",".dxr",".x32",".db",".jnt"}
ARCHIVE_EXT = {".zip",".rar",".7z"}
DTC_RE = re.compile(r"\b(?:P|B|C|U)\d{4}\b|\bDF\d{3,4}\b", re.I)
MAX_TEXT_CHARS = 30_000_000
CHUNK_CHARS = 6000
OVERLAP = 350
SHARD_TARGET = 18_000_000

def sha256_file(p):
    h=hashlib.sha256()
    with p.open("rb") as f:
        for b in iter(lambda:f.read(1024*1024), b""):
            h.update(b)
    return h.hexdigest()

def safe_dest(root, name):
    root=root.resolve()
    dest=(root / name).resolve()
    if root != dest and root not in dest.parents:
        raise ValueError("path traversal: "+name)
    return dest

def extract_zip(src, dst, report):
    dst.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(src) as z:
        for info in z.infolist():
            name=info.filename.replace("\\","/")
            if not name or name.endswith("/"):
                continue
            mode=(info.external_attr >> 16) & 0xF000
            if mode == 0xA000:
                report["extract_errors"].append({"file":name,"error":"symlink skipped"})
                continue
            try:
                out=safe_dest(dst,name)
            except Exception as e:
                report["extract_errors"].append({"file":name,"error":str(e)})
                continue
            out.parent.mkdir(parents=True, exist_ok=True)
            try:
                with z.open(info) as fi, out.open("wb") as fo:
                    shutil.copyfileobj(fi,fo,1024*1024)
            except Exception as e:
                report["extract_errors"].append({"file":name,"error":str(e)})

def extract_nested(root, report, max_depth=2):
    seen=set()
    for _depth in range(max_depth):
        archives=[p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in ARCHIVE_EXT and p not in seen]
        if not archives:
            break
        for p in archives:
            seen.add(p)
            rel=p.relative_to(root).as_posix()
            out=root / "__nested__" / re.sub(r"[^A-Za-z0-9._-]+","_",rel)[:180]
            out.mkdir(parents=True,exist_ok=True)
            try:
                subprocess.run(["7z","x","-y",str(p),"-o"+str(out)],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,timeout=1800)
                report["nested_archives"].append({"path":rel,"status":"extracted","to":out.relative_to(root).as_posix()})
            except Exception as e:
                report["nested_archives"].append({"path":rel,"status":"failed","error":str(e)})

def decode_bytes(data):
    for enc in ("utf-8","utf-8-sig","cp1252","latin1"):
        try:
            return data.decode(enc)
        except Exception:
            pass
    return data.decode("utf-8","ignore")

def clean_text(s):
    s=s.replace("\x00"," ")
    s=re.sub(r"[ \t]+"," ",s)
    s=re.sub(r"\n{4,}","\n\n\n",s)
    return s.strip()

def ocr_image(path):
    from PIL import Image, ImageOps
    import pytesseract
    with Image.open(path) as im:
        if getattr(im,"is_animated",False):
            try: im.seek(0)
            except Exception: pass
        im=ImageOps.exif_transpose(im.convert("RGB"))
        if max(im.size)>1800:
            ratio=1800.0/max(im.size)
            im=im.resize((max(1,int(im.width*ratio)),max(1,int(im.height*ratio))))
        im=ImageOps.autocontrast(ImageOps.grayscale(im))
        data=pytesseract.image_to_data(im,lang="por+eng",config="--psm 11",output_type=pytesseract.Output.DICT)
    words=[]; confs=[]
    for txt,cf in zip(data.get("text",[]),data.get("conf",[])):
        txt=(txt or "").strip()
        try: cf=float(cf)
        except Exception: cf=-1
        if txt:
            words.append(txt)
            if cf>=0: confs.append(cf)
    text=clean_text(" ".join(words))[:MAX_TEXT_CHARS]
    confidence=(sum(confs)/len(confs)/100.0) if confs else 0.0
    return text,max(0.0,min(1.0,confidence))

def legacy_doc_text(path):
    try:
        p=subprocess.run(["antiword",str(path)],stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=120,check=False)
        return clean_text(decode_bytes(p.stdout))[:MAX_TEXT_CHARS]
    except Exception:
        return ""

def binary_strings_text(path):
    try:
        p=subprocess.run(["strings","-a","-n","5",str(path)],stdout=subprocess.PIPE,stderr=subprocess.DEVNULL,timeout=180,check=False)
        return clean_text(decode_bytes(p.stdout[:50_000_000]))[:MAX_TEXT_CHARS]
    except Exception:
        return ""

def pdf_text(path):
    from pypdf import PdfReader
    r=PdfReader(str(path), strict=False)
    for i,page in enumerate(r.pages,1):
        try:
            t=clean_text(page.extract_text() or "")
        except Exception:
            t=""
        if t:
            yield {"page":i,"text":t[:MAX_TEXT_CHARS]}

def docx_text(path):
    from docx import Document
    d=Document(str(path))
    parts=[p.text for p in d.paragraphs if p.text]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return clean_text("\n".join(parts))[:MAX_TEXT_CHARS]

def pptx_text(path):
    from pptx import Presentation
    prs=Presentation(str(path))
    parts=[]
    for si,slide in enumerate(prs.slides,1):
        for shape in slide.shapes:
            if hasattr(shape,"text") and shape.text:
                parts.append(f"[slide {si}] "+shape.text)
    return clean_text("\n".join(parts))[:MAX_TEXT_CHARS]

def xlsx_text(path):
    import openpyxl
    wb=openpyxl.load_workbook(str(path),read_only=True,data_only=True)
    out=[]
    used=0
    for ws in wb.worksheets:
        out.append(f"[sheet {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            line=" | ".join("" if v is None else str(v) for v in row)
            if line.strip():
                out.append(line)
                used += len(line)+1
                if used>=MAX_TEXT_CHARS:
                    return clean_text("\n".join(out))[:MAX_TEXT_CHARS]
    return clean_text("\n".join(out))[:MAX_TEXT_CHARS]

def html_text(path):
    from bs4 import BeautifulSoup
    data=path.read_bytes()[:50_000_000]
    soup=BeautifulSoup(decode_bytes(data),"html.parser")
    return clean_text(soup.get_text("\n"))[:MAX_TEXT_CHARS]

def extract_text(path):
    ext=path.suffix.lower()
    if ext==".pdf":
        return "pages", list(pdf_text(path)), "pdf_text", 1.0, "source_extracted"
    if ext in IMAGE_EXT:
        t,conf=ocr_image(path)
        return "text", t, "ocr_tesseract", conf, "ocr_unverified"
    if ext==".docx":
        return "text", docx_text(path), "docx_text", 1.0, "source_extracted"
    if ext==".doc":
        return "text", legacy_doc_text(path), "antiword", 1.0, "source_extracted"
    if ext==".pptx":
        return "text", pptx_text(path), "pptx_text", 1.0, "source_extracted"
    if ext in {".xlsx",".xlsm"}:
        return "text", xlsx_text(path), "xlsx_text", 1.0, "source_extracted"
    if ext in {".html",".htm"}:
        return "text", html_text(path), "html_text", 1.0, "source_extracted"
    if ext in TEXT_EXT:
        return "text", clean_text(decode_bytes(path.read_bytes()[:50_000_000]))[:MAX_TEXT_CHARS], "text_decode", 1.0, "source_extracted"
    if ext in BINARY_TEXT_EXT:
        return "text", binary_strings_text(path), "binary_strings", 0.25, "unverified_strings"
    return "none", None, "none", 0.0, "metadata_only"

def chunks(text):
    if not text:
        return
    n=len(text)
    i=0
    while i<n:
        end=min(n,i+CHUNK_CHARS)
        yield text[i:end]
        if end==n:
            break
        i=max(i+1,end-OVERLAP)

class Shards:
    def __init__(self,outdir):
        self.outdir=outdir
        outdir.mkdir(parents=True,exist_ok=True)
        self.idx=0
        self.bytes=0
        self.fp=None
        self.path=None
        self.manifest=[]
    def rotate(self):
        if self.fp:
            self.fp.close()
            self.manifest[-1]["bytes"]=self.path.stat().st_size
        self.path=self.outdir/f"shard-{self.idx:04d}.jsonl.gz"
        self.fp=gzip.open(self.path,"wt",encoding="utf-8",compresslevel=6)
        self.manifest.append({"file":self.path.name,"records":0,"bytes":0})
        self.idx+=1
        self.bytes=0
    def add(self,obj):
        line=json.dumps(obj,ensure_ascii=False,separators=(",",":"))+"\n"
        if self.fp is None or self.bytes+len(line)>SHARD_TARGET:
            self.rotate()
        self.fp.write(line)
        self.bytes+=len(line)
        self.manifest[-1]["records"]+=1
    def close(self):
        if self.fp:
            self.fp.close()
            self.manifest[-1]["bytes"]=self.path.stat().st_size
            self.fp=None

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--archive",required=True)
    ap.add_argument("--extract-dir",required=True)
    ap.add_argument("--out",required=True)
    ap.add_argument("--nested-depth",type=int,default=2)
    args=ap.parse_args()
    archive=Path(args.archive)
    root=Path(args.extract_dir)
    out=Path(args.out)
    out.mkdir(parents=True,exist_ok=True)
    report={"archive":archive.name,"archive_bytes":archive.stat().st_size,"extract_errors":[],"nested_archives":[]}
    report["archive_sha256"]=sha256_file(archive)

    with zipfile.ZipFile(archive) as z:
        infos=[i for i in z.infolist() if not i.is_dir()]
        report["zip_entries"]=len(infos)
        report["zip_uncompressed_bytes"]=sum(i.file_size for i in infos)
        report["zip_largest_entry_bytes"]=max((i.file_size for i in infos),default=0)

    extract_zip(archive,root,report)
    extract_nested(root,report,args.nested_depth)

    files=[p for p in root.rglob("*") if p.is_file()]
    report["extracted_files"]=len(files)
    report["extracted_bytes"]=sum(p.stat().st_size for p in files)

    catalog=[]
    dtc_index={}
    ext_stats={}
    extractor_stats={"total_files":len(files),"indexed_files":0,"indexed_chunks":0,"unreadable_files":0,"ocr_files":0,"metadata_only_files":0,"unverified_string_files":0}
    shards=Shards(out/"brain")

    for n,p in enumerate(files,1):
        rel=p.relative_to(root).as_posix()
        size=p.stat().st_size
        ext=p.suffix.lower() or "(none)"
        ext_stats.setdefault(ext,{"files":0,"bytes":0})
        ext_stats[ext]["files"]+=1
        ext_stats[ext]["bytes"]+=size
        meta={"path":rel,"size":size,"ext":ext,"sha256":sha256_file(p)}
        kind="none"
        payload=None
        method="none"
        confidence=0.0
        truth_status="metadata_only"
        err=None
        try:
            kind,payload,method,confidence,truth_status=extract_text(p)
        except Exception as e:
            err=str(e)
            extractor_stats["unreadable_files"]+=1
        meta["indexed"]=bool(payload)
        meta["method"]=method
        meta["confidence"]=round(float(confidence or 0),4)
        meta["truth_status"]=truth_status
        if err:
            meta["extract_error"]=err[:500]
        catalog.append(meta)

        records=[]
        if method=="ocr_tesseract" and payload: extractor_stats["ocr_files"]+=1
        if truth_status=="unverified_strings" and payload: extractor_stats["unverified_string_files"]+=1
        if not payload: extractor_stats["metadata_only_files"]+=1
        if kind=="pages" and payload:
            for pg in payload:
                for ci,ch in enumerate(chunks(pg["text"])):
                    records.append({"path":rel,"page":pg["page"],"chunk":ci,"text":ch,"method":method,"confidence":confidence,"truth_status":truth_status})
        elif kind=="text" and payload:
            for ci,ch in enumerate(chunks(payload)):
                records.append({"path":rel,"chunk":ci,"text":ch,"method":method,"confidence":confidence,"truth_status":truth_status})

        if records:
            extractor_stats["indexed_files"]+=1

        for rec in records:
            rec["sha256"]=meta["sha256"]
            shards.add(rec)
            extractor_stats["indexed_chunks"]+=1
            for code in DTC_RE.findall(rec["text"]):
                code=code.upper()
                bucket=dtc_index.setdefault(code,[])
                if len(bucket)<250 and rel not in bucket:
                    bucket.append(rel)

        if n%100==0:
            print(f"indexed {n}/{len(files)} files",flush=True)

    shards.close()
    report["schema_version"]=2
    report["extractor"]=extractor_stats
    report["truth_policy"]={"source_extracted":"texto extraido diretamente sem correcao semantica","ocr_unverified":"OCR da imagem, nunca promovido sozinho a fato exato","unverified_strings":"strings recuperadas de binario proprietario, nunca usadas sozinhas como fato exato","metadata_only":"arquivo catalogado por caminho/hash/tamanho sem texto extraido"}
    report["extensions"]=ext_stats
    report["shards"]=shards.manifest

    (out/"manifest.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    with gzip.open(out/"catalog.jsonl.gz","wt",encoding="utf-8",compresslevel=6) as f:
        for row in catalog:
            f.write(json.dumps(row,ensure_ascii=False,separators=(",",":"))+"\n")
    with gzip.open(out/"source-records.jsonl.gz","wt",encoding="utf-8",compresslevel=6) as f:
        for row in catalog:
            f.write(json.dumps(row,ensure_ascii=False,separators=(",",":"))+"\n")
    (out/"dtc-index.json").write_text(json.dumps(dtc_index,ensure_ascii=False,separators=(",",":")),encoding="utf-8")
    (out/"README.md").write_text(
        "# thIAguinho Knowledge Branch\n\n"
        "Indice tecnico gerado automaticamente a partir do acervo CONHECIMENTO.ZIP.\n"
        "Os arquivos-fonte extraidos sao preservados em partes compactadas na Release knowledge-v1.\n"
        "O branch knowledge guarda o catalogo e os shards textuais pesquisaveis.\n",
        encoding="utf-8"
    )

    print(json.dumps({
        "archive_bytes":report["archive_bytes"],
        "zip_entries":report["zip_entries"],
        "zip_uncompressed_bytes":report["zip_uncompressed_bytes"],
        "extracted_files":report["extracted_files"],
        "extracted_bytes":report["extracted_bytes"],
        "extractor":report["extractor"]
    },indent=2))

if __name__=="__main__":
    main()
